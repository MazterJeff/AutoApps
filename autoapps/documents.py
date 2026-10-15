"""Render cover letters to PDF and DOCX for upload or download."""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path

from .models import Job, Profile


def _paragraphs(text: str) -> list[str]:
    paras, buf = [], []
    for line in text.splitlines():
        if line.strip():
            buf.append(line.strip())
        elif buf:
            paras.append(" ".join(buf))
            buf = []
    if buf:
        paras.append(" ".join(buf))
    return paras


def _contact_line(profile: Profile) -> str:
    return "  |  ".join(x for x in [profile.email, profile.phone, profile.location, profile.linkedin] if x)


def cover_letter_pdf(text: str, profile: Profile, job: Job) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from xml.sax.saxutils import escape

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=inch, rightMargin=inch, topMargin=0.9 * inch,
                            bottomMargin=0.9 * inch, title=f"Cover Letter - {job.company}", author=profile.full_name)
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["Normal"], fontName="Helvetica", fontSize=11, leading=15, spaceAfter=10)
    name = ParagraphStyle("name", parent=body, fontName="Helvetica-Bold", fontSize=15, leading=19, spaceAfter=2)
    small = ParagraphStyle("small", parent=body, fontSize=9.5, leading=12, textColor="#444444", spaceAfter=18)
    story = []
    if profile.full_name:
        story.append(Paragraph(escape(profile.full_name), name))
    if _contact_line(profile):
        story.append(Paragraph(escape(_contact_line(profile)), small))
    story.append(Paragraph(date.today().strftime("%B %d, %Y").replace(" 0", " "), body))
    story.append(Spacer(1, 4))
    for para in _paragraphs(text):
        story.append(Paragraph(escape(para), body))
    doc.build(story)
    return buf.getvalue()


def cover_letter_docx(text: str, profile: Profile, job: Job) -> bytes:
    import docx
    from docx.shared import Pt

    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    if profile.full_name:
        run = d.add_paragraph().add_run(profile.full_name)
        run.bold = True
        run.font.size = Pt(15)
    if _contact_line(profile):
        d.add_paragraph(_contact_line(profile))
    d.add_paragraph(date.today().strftime("%B %d, %Y").replace(" 0", " "))
    for para in _paragraphs(text):
        d.add_paragraph(para)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def write_cover_letter_pdf(path: Path, text: str, profile: Profile, job: Job) -> Path:
    path.write_bytes(cover_letter_pdf(text, profile, job))
    return path
