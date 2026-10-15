"""Extract text from an uploaded resume."""

from __future__ import annotations

import io
from pathlib import Path

ALLOWED = {".pdf", ".docx", ".txt", ".md"}


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    if ext == ".docx":
        import docx

        doc = docx.Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts).strip()
    if ext in {".txt", ".md"}:
        return data.decode("utf-8", errors="replace").strip()
    raise ValueError(f"Unsupported resume type {ext!r}; use one of {', '.join(sorted(ALLOWED))}")
